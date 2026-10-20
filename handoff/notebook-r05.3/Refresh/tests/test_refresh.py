from __future__ import annotations

import csv
import importlib.util
import io
import os
import sys
from pathlib import Path

import pytest
from hypothesis import given, strategies as st
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

SOURCE = Path(__file__).resolve().parents[1] / "source"
sys.path.insert(0, str(SOURCE))
import fast_refresh as fast
import model_refresh as models
import reporting
import catalog

spec = importlib.util.spec_from_file_location("fast_reader", SOURCE / "reader.py")
reader = importlib.util.module_from_spec(spec)
spec.loader.exec_module(reader)
NAME = "UK_Example_Bridge_Jul 26_EF-Jun 26_IMS-May 26_ADS.xlsm"


def workbook(value=100.0, profile="Ext"):
    book = Workbook()
    page = book.active
    page.title = "Example"
    metadata = {
        "Country": "UK",
        "GrowthDriver": "Example",
        "Cycle": "2026-07",
        "BridgeRevision": "ADS_Rev001",
        "SourceProfile": profile,
        "PublicationStatus": "DEVELOPMENT",
        "PublicationOrigin": "TEST_ONLY",
        "ExFactoryActualsThrough": "2026-06-01",
        "IMSActualsThrough": "2026-05-01",
        "PreviousPublication": "",
    }
    for k, v in metadata.items():
        book.defined_names.add(DefinedName("Bridge_" + k, attr_text='"' + v + '"'))
    headers = [*reader.IDS, "2026-Apr", "2026-May", "2026-Jun"]
    for col, text in enumerate(headers, 1):
        page.cell(65, col, text)
    rows = [
        ("", "", "Example", "", "", "Model", "Working days", 20, None, 0),
        (
            "Community",
            "",
            "Example",
            "Abbott",
            "Example Abbott",
            "Pricing",
            "Product Selling units",
            value,
            0,
            None,
        ),
        (
            "Community",
            "",
            "Example",
            "Other",
            "Example Other",
            "Pricing",
            "Product Selling units",
            50,
            None,
            None,
        ),
        ("", "", "Example", "", "", "New section", "New measure", "0", True, None),
        ("", "", "Example", "", "", "Pricing", "Residual Product", None, None, None),
        ("", "", "Example", "", "", "Pricing", "Status", "PASS", "CHECK", "NO DATA"),
    ]
    for number, row in enumerate(rows, 66):
        for col, val in enumerate(row, 1):
            page.cell(number, col, val)
    page.add_table(Table(displayName="SectorModel01", ref="A65:J71"))
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    "folder",
    [
        "",
        "My work/",
        "2026/10 Oct 26/Bridge Files/Archive/",
        "Someone/New Folder/Versions/",
        "V18/",
        "2025/not the cycle/",
    ],
)
@pytest.mark.parametrize(
    "suffix", ["ADS", "ADS_Rev001", "Indi_001", "my working forecast V2", "V18_Ext"]
)
def test_folder_independent_versions(folder, suffix):
    name = NAME.replace("ADS.xlsm", suffix + ".xlsm")
    assert fast.filename_candidate(name)
    assert __import__("fnmatch").fnmatchcase(name, fast.FILE_GLOB)
    parsed = reader.BridgeReader(workbook(), folder + name)
    result = list(parsed.rows())
    assert parsed.metadata["DemandCycle"] == "2026-07"
    assert parsed.metadata["SourceProfile"] == "Ext"
    assert parsed.metadata["PublicationStatus"] == "DEVELOPMENT"
    assert len(result) == 6
    assert parsed.observation_count == 18


@given(
    st.text(
        alphabet=st.characters(
            blacklist_categories=("Cs",), blacklist_characters="/\\\x00\n\r"
        ),
        min_size=1,
        max_size=50,
    )
)
def test_any_nonempty_version_suffix(suffix):
    assert fast.filename_candidate(NAME.replace("ADS.xlsm", suffix + ".xlsx"))


@pytest.mark.parametrize(
    "name",
    [
        "~$" + NAME,
        NAME.replace(".xlsm", ".xlsb"),
        NAME.replace(".xlsm", ".xls"),
        "DemandMaster.xlsx",
        NAME.replace("_EF-Jun 26", ""),
        NAME + ".bak",
    ],
)
def test_outside_convention_names_do_not_enter_reader(name):
    assert not fast.filename_candidate(name)


def test_local_and_ext_semantics_and_types():
    for profile in ["Local", "Ext"]:
        r = reader.BridgeReader(workbook(profile=profile), "arbitrary/" + NAME)
        rows = list(r.rows())
        assert r.metadata["SourceProfile"] == profile
        p = next(x for x in rows if x["ids"][4] == "Example Abbott")
        assert [m[4] for m in p["months"]] == [100.0, 0.0, None]
        special = next(x for x in rows if x["ids"][6] == "New measure")
        assert [m[2] for m in special["months"]] == ["text", "bool", "blank"]
    with pytest.raises(ValueError, match="cycle/profile"):
        list(reader.BridgeReader(workbook(), NAME.replace("Jul 26", "Aug 26")).rows())


def test_copied_paths_have_distinct_report_identity():
    sha = "a" * 64
    assert fast.report_identity(sha, "original/" + NAME) == fast.report_identity("b" * 64, "ORIGINAL/" + NAME)
    assert fast.report_identity(sha, "my work/" + NAME, "original/" + NAME) != sha
    assert fast.report_identity(
        sha, "my work/" + NAME, "original/" + NAME
    ) == fast.report_identity(sha, "my work/" + NAME, "original/" + NAME)


class Response:
    def __init__(self, status=200, body=None, headers=None):
        self.status_code = status
        self.body = body or {}
        self.headers = headers or {}
        self.text = str(self.body)

    def json(self):
        return self.body


class HTTP:
    def __init__(self, states=None, post_status=202):
        self.states = list(states or ["Completed"])
        self.posts = []
        self.post_status = post_status

    def post(self, url, **kw):
        self.posts.append(kw)
        return Response(
            self.post_status,
            headers={"Location": url + "/abc", "x-ms-request-id": "abc"},
        )

    def get(self, url, **kw):
        return Response(body={"status": self.states.pop(0)})


def api_config():
    return {
        "WORKSPACE_ID": "f5cb5742-9a7c-4337-a4c0-952706dddf0b",
        "SEMANTIC_MODEL_ID": "e7af2967-7725-4108-a170-9e490d4e35bb",
    }


def test_model_refresh_waits_for_completion_and_never_notifies():
    http = HTTP(["InProgress", "Completed"])
    r = models.refresh_model(
        api_config(), lambda audience: "test-token", http=http, sleep=lambda n: None
    )
    assert r["Status"] == "COMPLETED"
    assert http.posts[0]["json"] == {
        "type": "automatic",
        "commitMode": "transactional",
        "retryCount": 1,
    }
    assert "notifyOption" not in http.posts[0]["json"]


def test_model_refresh_failed_and_rejected_are_not_success():
    with pytest.raises(RuntimeError, match="Failed"):
        models.refresh_model(
            api_config(), lambda audience: "test-token", http=HTTP(["Failed"])
        )
    with pytest.raises(RuntimeError, match="not accepted"):
        models.refresh_model(
            api_config(), lambda audience: "test-token", http=HTTP(post_status=403)
        )


def test_model_refresh_unknown_location_is_not_completed():
    h = HTTP()
    h.post = lambda url, **kw: Response(
        202, headers={"Location": "https://example.com/invalid"}
    )
    assert (
        models.refresh_model(api_config(), lambda a: "test-token", http=h)["Status"]
        == "ACCEPTED_UNCONFIRMED"
    )


@pytest.fixture(scope="session")
def spark(tmp_path_factory):
    import jdk4py
    from pyspark.sql import SparkSession

    os.environ["JAVA_HOME"] = str(jdk4py.JAVA_HOME)
    os.environ["PYSPARK_PYTHON"] = sys.executable
    os.environ["SPARK_LOCAL_IP"] = "127.0.0.1"
    root = tmp_path_factory.mktemp("spark")
    jars = ",".join(str(p) for p in Path("/tmp/forecastmaster-fast-jars").glob("*.jar"))
    session = (
        SparkSession.builder.master("local[2]")
        .appName("ForecastMaster synthetic acceptance")
        .config("spark.jars", jars)
        .config("spark.ui.enabled", "false")
        .config("spark.sql.extensions", "io.delta.sql.DeltaSparkSessionExtension")
        .config(
            "spark.sql.catalog.spark_catalog",
            "org.apache.spark.sql.delta.catalog.DeltaCatalog",
        )
        .config("spark.sql.warehouse.dir", str(root / "warehouse"))
        .config("spark.sql.shuffle.partitions", "2")
        .config("spark.default.parallelism", "2")
        .config("spark.databricks.delta.snapshotPartitions", "2")
        .getOrCreate()
    )
    session.sparkContext.setLogLevel("ERROR")
    yield session
    session.stop()


@pytest.mark.integration
def test_real_delta_retry_alias_projection_nochange_and_edit(
    spark, tmp_path, monkeypatch
):
    from pyspark.sql import functions as F
    from test_legacy_metadata import metadata_variant

    root = tmp_path / "Demand"
    one = root / "My Working Folder" / NAME
    one.parent.mkdir(parents=True)
    one.write_bytes(metadata_variant(workbook(), {
        "SourceProfile": "V15_Ext", "BridgeRevision": "V17_Ext", "PublicationOrigin": None,
    }))
    alias = (
        root
        / "Someone Else"
        / NAME.replace("ADS.xlsm", "ADS_Rev001.xlsm").replace("_Bridge_", "_bridge_")
    )
    alias.parent.mkdir(parents=True)
    alias.write_bytes(one.read_bytes())
    (root / "DemandMaster.xlsx").write_bytes(b"not a bridge")
    config = {
        "SOURCE_ROOT": str(root),
        "TABLE_PREFIX": "fast_test",
        "ALLOWED_PROFILES": ["Local", "Ext"],
        "REPORT_PREFIX": "fast_report",
        "REFRESH_MODE": "ALL",
        "REPORT_RULES_PATH": str(
            tmp_path / "Files/ForecastMasterControls/ReportMeasureRules.csv"
        ),
    }
    r = fast.FastRefresh(spark, vars(reader), config)
    actual = r.upsert

    def fail_midway(name, frame, keys, **kwargs):
        if name.endswith("_calendar"):
            raise RuntimeError("injected interruption before source receipt")
        return actual(name, frame, keys, **kwargs)

    r.upsert = fail_midway
    r.ingest()
    assert r.failed_files == 2
    assert not spark.catalog.tableExists("fast_test_source_revisions")
    assert spark.table("fast_test_history").count() == 18
    assert (
        spark.table("fast_test_history").select("ObservationID").distinct().count()
        == 18
    )
    retry = fast.FastRefresh(spark, vars(reader), config)
    retry.ingest()
    assert retry.failed_files == 0
    assert len(retry.completed) == 1
    assert spark.table("fast_test_history").count() == 18
    assert spark.table("fast_test_source_revisions").count() == 1
    assert len(retry.report_sources()) == 2
    assert retry.writes["fast_test_ingestion_events"] == 1
    assert retry.writes["fast_test_history"] == 1
    # Rules are the actual seven reviewed report definitions, with synthetic source values.
    rules = (SOURCE.parent / "ReportMeasureRules.csv").read_text(encoding="utf-8-sig")
    monkeypatch.chdir(tmp_path)
    controls = tmp_path / "Files" / "ForecastMasterControls"
    controls.mkdir(parents=True)
    (controls / "ReportMeasureRules.csv").write_text(rules)
    retry.config["MEASURE_CATALOG_FILE"] = str(controls / "MeasureCatalog.csv")
    (controls / "MeasureCatalog.csv").write_text(
        "SourceSection,SourceMeasure,ReportingEnabled\nPricing,Product Selling units,true\n"
    )
    assert catalog.update_catalog(retry)
    assert (
        spark.table("fast_test_measure_assignments")
        .where("CatalogState = 'MAPPED'")
        .count()
        == 1
    )
    # Simulate an installed pre-R04.1 gallery and its three-column receipts.
    from datetime import date
    spark.createDataFrame([("obsolete-import", "Community", "", "Example", date(2026,4,1), "product-units", 999.0, 1, 0)], "SourceRevisionID string, Channel string, SubChannel string, Sector string, Month date, ComponentKey string, Amount double, NumericCells long, MissingCells long").write.format("delta").mode("errorifexists").saveAsTable("fast_report_facts")
    spark.createDataFrame([("obsolete-import", "old-rules", "2026-01-01")], "SourceRevisionID string, RuleHash string, ProjectedAtUtc string").write.format("delta").mode("errorifexists").saveAsTable("fast_report_projection_receipts")
    assert reporting.update_reporting(retry, rules)
    assert "ContentRevisionID" in spark.table("fast_report_projection_receipts").columns
    assert spark.table("fast_report_facts").where("SourceRevisionID = 'obsolete-import'").count() == 0
    assert spark.table("fast_report_files").count() == 2
    facts = spark.table("fast_report_facts").where(
        F.col("ComponentKey") == "product-units"
    )
    april = facts.where(F.col("Month") == F.lit("2026-04-01")).collect()
    may = facts.where(F.col("Month") == F.lit("2026-05-01")).collect()
    june = facts.where(F.col("Month") == F.lit("2026-06-01")).collect()
    assert len(april) == 2 and all(
        x["Amount"] == 150.0 and x["NumericCells"] == 2 for x in april
    )
    assert len(may) == 2 and all(
        x["Amount"] == 0.0 and x["MissingCells"] == 1 for x in may
    )
    assert len(june) == 2 and all(
        x["Amount"] is None and x["NumericCells"] == 0 and x["MissingCells"] == 2
        for x in june
    )
    again = fast.FastRefresh(spark, vars(reader), config)
    again.ingest()
    assert all(
        e["Status"] in {"METADATA_UNCHANGED", "OUTSIDE_FILENAME_PATTERN"}
        for e in again.events
    )
    again.config["MEASURE_CATALOG_FILE"] = retry.config["MEASURE_CATALOG_FILE"]
    assert not catalog.update_catalog(again)
    assert not reporting.update_reporting(again, rules)
    assert set(again.writes) == {"fast_test_ingestion_events"}
    first_ids = {r["SourceRelativePath"]: r["SourceRevisionID"] for r in retry.report_sources()}
    one.write_bytes(workbook(value=120.0))
    edit = fast.FastRefresh(
        spark,
        vars(reader),
        {
            **config,
            "REFRESH_MODE": "FILES",
            "SELECTED_RELATIVE_PATHS": ["My Working Folder/" + NAME],
        },
    )
    edit.ingest()
    assert edit.failed_files == 0 and len(edit.completed) == 2
    assert spark.table("fast_test_history").count() == 36
    assert len(edit.report_sources()) == 2
    assert reporting.update_reporting(edit, rules)
    assert spark.table("fast_report_files").count() == 2
    values = {
        x[0]
        for x in spark.table("fast_report_facts")
        .where(
            (F.col("ComponentKey") == "product-units")
            & (F.col("Month") == F.lit("2026-04-01"))
        )
        .select("Amount")
        .collect()
    }
    assert values == {150.0, 170.0}

    assert first_ids == {r["SourceRelativePath"]: r["SourceRevisionID"] for r in edit.report_sources()}
    assert spark.table("fast_report_facts").count() == 12
    edit.compact_current_sources()
    # The old bytes are still used by the separately named alias, so keep them.
    assert spark.table("fast_test_history").count() == 36
    one.write_bytes(workbook(value=140.0))
    newer = fast.FastRefresh(spark, vars(reader), edit.config)
    newer.ingest()
    assert reporting.update_reporting(newer, rules)
    newer.compact_current_sources()
    assert spark.table("fast_test_history").count() == 36
    assert spark.table("fast_test_source_revisions").count() == 2
    assert spark.table("fast_report_files").count() == 2
    assert spark.table("fast_report_facts").count() == 12
    assert {r[0] for r in spark.table("fast_report_facts").where("Month = '2026-04-01' AND ComponentKey = 'product-units'").select("Amount").collect()} == {150.0, 190.0}
    # Blank and zero overwrite old numbers instead of leaving stale MERGE rows.
    one.write_bytes(workbook(value=0.0))
    zero = fast.FastRefresh(spark, vars(reader), edit.config)
    zero.ingest(); reporting.update_reporting(zero, rules); zero.compact_current_sources()
    assert spark.table("fast_test_history").count() == 36
    assert spark.table("fast_report_facts").count() == 12
    assert {r[0] for r in spark.table("fast_report_facts").where("Month = '2026-04-01' AND ComponentKey = 'product-units'").select("Amount").collect()} == {50.0, 150.0}

    # A failed replacement retains the previous valid data and pointer.
    current_hashes = {r[0] for r in spark.table("fast_test_file_state").select("SourceRevisionID").collect()}
    one.write_bytes(b"incomplete upload")
    broken = fast.FastRefresh(spark, vars(reader), edit.config)
    broken.ingest()
    assert broken.failed_files == 1
    assert {r[0] for r in spark.table("fast_test_file_state").select("SourceRevisionID").collect()} == current_hashes
    assert not reporting.update_reporting(broken, rules)
    assert spark.table("fast_test_history").count() == 36

    # Replacing a file also removes dimension rows it no longer contains.
    from openpyxl import load_workbook
    book=load_workbook(io.BytesIO(workbook(value=10.0)))
    book.active.cell(67,1,"Out of Pocket");book.active.cell(68,1,"Out of Pocket")
    blob=io.BytesIO();book.save(blob);one.write_bytes(blob.getvalue())
    changed_dimension=fast.FastRefresh(spark, vars(reader), edit.config)
    changed_dimension.ingest();reporting.update_reporting(changed_dimension,rules);changed_dimension.compact_current_sources()
    one_id=first_ids["My Working Folder/"+NAME]
    assert spark.table("fast_report_facts").where((F.col("SourceRevisionID")==one_id)&(F.col("Channel")=="Community")).count() == 0
    assert spark.table("fast_report_facts").where((F.col("SourceRevisionID")==one_id)&(F.col("Channel")=="Out of Pocket")).count() == 6
    assert spark.table("fast_test_history").count() == 36
    assert spark.table("fast_report_facts").count() == 12

    # A report-definition failure must release its cached projection and preserve old facts.
    before = spark.table("fast_report_facts").count()
    bad_rules = rules.replace(
        "Pricing,Product Selling units", "Unknown source,Unknown units"
    )
    (controls / "ReportMeasureRules.csv").write_text(bad_rules)
    from pyspark.sql import DataFrame

    cached = []
    original_persist = DataFrame.persist

    def track_persist(frame, *args, **kwargs):
        result = original_persist(frame, *args, **kwargs)
        cached.append(result)
        return result

    monkeypatch.setattr(DataFrame, "persist", track_persist)
    with pytest.raises(ValueError, match="No fact components"):
        reporting.update_reporting(changed_dimension, bad_rules)
    assert spark.table("fast_report_facts").count() == before
    assert len(cached) == 1 and not cached[0].is_cached

    # A complete inventory retires a deleted scenario and reclaims its old data.
    (controls / "ReportMeasureRules.csv").write_text(rules)
    alias.unlink()
    retired=fast.FastRefresh(spark, vars(reader), config)
    retired.ingest();catalog.update_catalog(retired);reporting.update_reporting(retired,rules);retired.compact_current_sources()
    assert spark.table("fast_test_file_state").count() == 1
    assert spark.table("fast_report_files").count() == 1
    assert spark.table("fast_test_source_revisions").count() == 1
    assert spark.table("fast_test_history").count() == 18
    assert spark.table("fast_report_facts").count() == 6
    assert retired.report_sources()[0]["SourceRevisionID"] == one_id
